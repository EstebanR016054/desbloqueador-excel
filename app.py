import streamlit as st
import openpyxl
import io

st.set_page_config(page_title="Desbloqueador de Excel", page_icon="🔓")

st.title("🔓 Desbloqueador de Hojas de Excel")
st.write("Sube tu archivo Excel. El sistema buscará las hojas que estén visibles y tengan datos en la celda Z3, y las desprotegerá automáticamente.")

# 1. Espacio para que el usuario cargue el archivo
archivo_subido = st.file_uploader("Elige un archivo Excel (.xlsx)", type=["xlsx"])

if archivo_subido is not None:
    st.info("Procesando archivo... por favor espera.")
    
    try:
        # 2. Leer el Excel
        wb = openpyxl.load_workbook(archivo_subido)
        hojas_modificadas = 0
        
        # 3. Recorrer cada hoja y aplicar tus filtros
        for nombre_hoja in wb.sheetnames:
            ws = wb[nombre_hoja]
            
            # Filtro A: ¿La hoja está visible?
            if ws.sheet_state != 'visible':
                continue # Salta a la siguiente
                
            # Filtro B: ¿Tiene datos en la celda Z3?
            valor_z3 = ws['Z3'].value
            if valor_z3 is None or str(valor_z3).strip() == "":
                continue # Salta a la siguiente
                
            # Si cumple las condiciones, quitamos la protección (sin necesidad de contraseña)
            if ws.protection.sheet:
                ws.protection.sheet = False
                hojas_modificadas += 1
                st.write(f"✅ Hoja desprotegida: **{nombre_hoja}**")
            
        # 4. Preparar el archivo modificado para descargarlo
        if hojas_modificadas > 0:
            salida = io.BytesIO()
            wb.save(salida)
            salida.seek(0)
            
            st.success(f"¡Proceso terminado! Se desprotegieron {hojas_modificadas} hoja(s).")
            
            # 5. Botón para descargar el archivo nuevo
            st.download_button(
                label="Descargar Excel Desprotegido",
                data=salida,
                file_name="Desprotegido_" + archivo_subido.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            st.warning("No se encontró ninguna hoja que estuviera protegida, que fuera visible y que tuviera datos en Z3.")
            
    except Exception as e:
        st.error(f"Hubo un error al procesar el archivo: {e}")